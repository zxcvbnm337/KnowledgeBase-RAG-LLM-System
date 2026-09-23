"""多格式文档解析：把上传的文件统一转成「纯文本 + 元数据」。

为什么单独抽一层：
    入库前的「解析」是 RAG 全流程的第一环，也是最容易悄悄劣化的一环——
    PDF 抽不出文字、Excel 只读了第一个 sheet，都不会报错，只会让检索质量下降。
    把解析集中在这里，既方便单测，也方便未来换用更专业的解析器（如 MinerU / Unstructured）。

依赖策略：
    pypdf / python-docx / openpyxl 都是**延迟导入**的可选依赖，
    缺失时给出明确提示而不是抛一长串 traceback。
"""

from __future__ import annotations

import io
import os

# 支持的扩展名（全部小写）
SUPPORTED_EXTENSIONS = (".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx")

# 前端 file_uploader 的 type 参数（不带点）
UPLOADER_TYPES = [ext.lstrip(".") for ext in SUPPORTED_EXTENSIONS]

# 纯文本类：直接解码，无需第三方解析库
TEXT_EXTENSIONS = (".txt", ".md", ".csv")

_DECODE_ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "gb18030")


class UnsupportedFileType(ValueError):
    """扩展名不在支持范围内。"""


class ParseError(RuntimeError):
    """文件受支持，但解析失败（内容损坏、缺少可选依赖等）。"""


def _decode(data: bytes, filename: str) -> str:
    """按常见中文编码依次尝试解码。

    只用 utf-8 会在 GBK 文件上直接崩——这是国内场景下最高频的解析失败原因。
    """
    for enc in _DECODE_ENCODINGS:
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ParseError(f"{filename}：尝试 {'/'.join(_DECODE_ENCODINGS)} 均解码失败，"
                     "请确认文件编码或另存为 UTF-8。")


def _parse_pdf(data: bytes, filename: str) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader  # 延迟导入
    except ImportError as e:
        raise ParseError(f"{filename}：解析 PDF 需要 pypdf，请先 pip install pypdf（{e}）")

    reader = PdfReader(io.BytesIO(data))
    pages = []
    empty_pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[第 {i} 页]\n{text}")
        else:
            empty_pages.append(i)

    text = "\n\n".join(pages)
    if not text:
        raise ParseError(f"{filename}：未能抽取到任何文字，可能是扫描件（图片型 PDF），"
                         "需要先做 OCR。")
    return text, {"source_type": "pdf", "pages": len(reader.pages),
                  "empty_pages": empty_pages}


def _parse_docx(data: bytes, filename: str) -> tuple[str, dict]:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise ParseError(f"{filename}：解析 docx 需要 python-docx，请先 pip install python-docx（{e}）")

    document = docx.Document(io.BytesIO(data))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # 表格同样承载业务信息，之前只读段落会把表格内容整体丢掉
    table_count = 0
    for table in document.tables:
        table_count += 1
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    text = "\n".join(blocks)
    if not text:
        raise ParseError(f"{filename}：文档内没有任何可提取的文字。")
    return text, {"source_type": "docx", "paragraphs": len(blocks),
                  "tables": table_count}


def _parse_xlsx(data: bytes, filename: str) -> tuple[str, dict]:
    try:
        import openpyxl
    except ImportError as e:
        raise ParseError(f"{filename}：解析 xlsx 需要 openpyxl，请先 pip install openpyxl（{e}）")

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    total_rows = 0
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            total_rows += len(rows)
            sections.append(f"[工作表：{ws.title}]\n" + "\n".join(rows))
    wb.close()

    if not sections:
        raise ParseError(f"{filename}：工作簿中没有任何非空单元格。")
    return "\n\n".join(sections), {"source_type": "xlsx",
                                   "sheets": len(wb.sheetnames),
                                   "rows": total_rows}


def load_document(filename: str, data: bytes) -> tuple[str, dict]:
    """把上传文件解析为 `(纯文本, 元数据)`。

    :raises UnsupportedFileType: 扩展名不在 SUPPORTED_EXTENSIONS 中
    :raises ParseError: 解析失败（含缺少可选依赖）
    """
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"暂不支持 {ext or '未知'} 格式。目前支持："
            + "、".join(e.lstrip('.') for e in SUPPORTED_EXTENSIONS)
        )

    if ext in TEXT_EXTENSIONS:
        text, meta = _decode(data, filename), {"source_type": ext.lstrip(".")}
    else:
        # 各格式解析器会抛出五花八门的底层异常（BadZipFile、PackageNotFoundError、
        # PdfReadError…）。统一收敛成 ParseError，调用方只需要 catch 一种异常，
        # 用户也不会看到一屏 traceback。
        try:
            if ext == ".pdf":
                text, meta = _parse_pdf(data, filename)
            elif ext == ".docx":
                text, meta = _parse_docx(data, filename)
            else:  # .xlsx
                text, meta = _parse_xlsx(data, filename)
        except ParseError:
            raise
        except Exception as e:
            raise ParseError(f"{filename}：解析失败，文件可能已损坏或格式不符"
                             f"（{type(e).__name__}: {e}）") from e

    text = text.strip()
    if not text:
        raise ParseError(f"{filename}：解析后内容为空。")

    meta["chars"] = len(text)
    meta["original_name"] = filename
    return text, meta


if __name__ == "__main__":
    demo = "这是一段用于验证解析链路的测试文本。".encode("utf-8")
    print(load_document("demo.txt", demo))
